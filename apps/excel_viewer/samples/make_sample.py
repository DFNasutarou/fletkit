"""動作確認用のサンプル Excel を作る。  python make_sample.py [出力先]"""

import datetime as dt
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment


def make(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "計測ログ"
    ws.append(["時刻", "車速[km/h]", "エンジン回転数[rpm]", "CAN ID", "状態"])
    rnd = random.Random(1)
    t0 = dt.datetime(2026, 10, 28, 9, 0, 0)
    for i in range(1500):
        ws.append([
            t0 + dt.timedelta(seconds=i),
            round(40 + 20 * rnd.random(), 1),
            rnd.randint(1500, 3200),
            f"0x{rnd.choice([0x1A0, 0x2B4, 0x3C8]):03X}",
            rnd.choice(["正常", "正常", "正常", "警告"]),
        ])
    ws["A2"].comment = Comment("コメントは表示されません", "sample")

    s2 = wb.create_sheet("集計")
    s2.append(["項目", "値", "備考"])
    s2.append(["平均車速", 50.2, None])
    s2.append(["最大回転数", 3199, "結合セルの例→"])
    s2.merge_cells("C3:D4")
    s2.append(["記録日", dt.date(2026, 10, 28)])
    s2.append(["計算式", "=B2*2"])

    hidden = wb.create_sheet("設定値")
    hidden.sheet_state = "hidden"
    hidden.append(["しきい値", 3000])

    wb.create_sheet("空のシート")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("sample.xlsx")
    print(make(out))
