"""Excel ファイルの読み込み（Flet・fletkit に依存しない部分）。

セルの値だけを読む。図形・コメント・書式・画像・グラフは扱わない。
- 数式は最後に保存された計算結果（data_only=True）。保存されていなければ空欄。
- 結合セルは左上のセルにだけ値が入り、残りは空欄。
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

SUPPORTED = (".xlsx", ".xlsm")
ROW_COLUMN = "行"


class ExcelLoadError(Exception):
    """利用者に見せるメッセージを持つ読み込みエラー。"""


@dataclass
class SheetInfo:
    name: str
    hidden: bool


@dataclass
class SheetValues:
    """シートのセル値（末尾の空行・空列は除く）。rows[i] は Excel の i+1 行目。"""

    name: str
    rows: list[list[Any]]

    @property
    def size(self) -> tuple[int, int]:
        return len(self.rows), (len(self.rows[0]) if self.rows else 0)

    def table(self, header_first_row: bool) -> tuple[list[str], list[list[Any]]]:
        """表示用の (列名, 行) を返す。先頭列は Excel の行番号。"""
        n_rows, n_cols = self.size
        letters = [get_column_letter(i + 1) for i in range(n_cols)]
        if n_rows == 0:
            return [], []
        if header_first_row:
            head = self.rows[0]
            names = [str(v).strip() if v not in (None, "") else letters[i] for i, v in enumerate(head)]
            body = self.rows[1:]
            start = 2
        else:
            names = letters
            body = self.rows
            start = 1
        return [ROW_COLUMN, *names], [[start + i, *r] for i, r in enumerate(body)]


class ExcelBook:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        if self.path.suffix.lower() not in SUPPORTED:
            raise ExcelLoadError(f"対応していない形式です：{self.path.suffix}（.xlsx / .xlsm のみ）")
        try:
            # 結合セルの情報が必要なので read_only は使わない
            self._wb = openpyxl.load_workbook(self.path, data_only=True)
        except FileNotFoundError:
            raise ExcelLoadError(f"ファイルが見つかりません：{self.path}") from None
        except PermissionError:
            raise ExcelLoadError(f"ファイルを開けません（他のアプリで使用中の可能性があります）：{self.path.name}") from None
        except (InvalidFileException, zipfile.BadZipFile, KeyError) as ex:
            raise ExcelLoadError(f"Excel ファイルとして読み込めません（パスワード付き・破損の可能性）：{self.path.name}\n{ex}") from None
        self._cache: dict[str, SheetValues] = {}

    def sheets(self, include_hidden: bool = False) -> list[SheetInfo]:
        out = []
        for ws in self._wb.worksheets:
            hidden = ws.sheet_state != "visible"
            if hidden and not include_hidden:
                continue
            out.append(SheetInfo(ws.title, hidden))
        return out

    def values(self, name: str) -> SheetValues:
        """シートのセル値を返す（初回だけ変換し、以後はキャッシュ）。"""
        cached = self._cache.get(name)
        if cached is not None:
            return cached
        ws = self._wb[name]
        rows = [list(r) for r in ws.iter_rows(min_row=1, min_col=1, values_only=True)]
        rows = _trim(rows)
        sv = SheetValues(name, rows)
        self._cache[name] = sv
        return sv

    def close(self) -> None:
        try:
            self._wb.close()
        except Exception:
            pass


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _trim(rows: list[list[Any]]) -> list[list[Any]]:
    """末尾の空行と、全行で空の末尾列を取り除き、各行の列数をそろえる。"""
    last_row: Optional[int] = None
    last_col = -1
    for i, r in enumerate(rows):
        for j in range(len(r) - 1, -1, -1):
            if not _blank(r[j]):
                last_row = i
                last_col = max(last_col, j)
                break
    if last_row is None:
        return []
    width = last_col + 1
    out = []
    for r in rows[: last_row + 1]:
        r = r[:width]
        out.append(r + [None] * (width - len(r)))
    return out
